"""C3 deliverable: the file quality metrics document and scores workbook (problem statement deliverable 3).

Deliverable 3 asks for "the metrics used to grade PTD files, how each is calculated, and why"; C3 asks to
"score every file you use and justify each metric". This script writes both from the C3 outputs, so the
numbers in the document always match the data:

  deliverables/03_file_quality_metrics.md   method, every metric (calculation, why, basis), scores of
                                            every file, the main problems per dataset
  deliverables/03_file_quality_scores.xlsx  the same definitions and every score, one sheet per table

Inputs: output/c3/ (c3_metrics.py, c3_index_metrics.py, c3_diagnose.py).

Usage:
    python pipeline/c3_report.py
"""
import argparse
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from c3_index_metrics import GROUPS as INDEX_GROUPS, INDEX_SHARE  # noqa: E402
from c3_metrics import CYCLE_DAYS, DECAY_DAYS, EXTREME, FENCE, GROUPS as FILE_GROUPS, MIN_PRICES  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
C3 = os.path.join(ROOT, "output", "c3")
OUT = os.path.join(ROOT, "deliverables")
DISPLAY = {"ambetter": "Ambetter", "healthnet": "Health Net", "fidelis": "Fidelis", "qualchoice": "QualChoice"}
GROUP_NAME = {"gate": "Gate", "timeliness": "Timeliness", "compliance": "Compliance", "plausibility": "Plausibility",
              "usability": "Usability", "consistency": "Consistency", "reported": "Reported only"}

RULE = "45 CFR §147.212 (Oct 2025)"
# id: (how it is calculated, why it matters, basis)
DEF = {
    "M1": ("GET on the URL the index lists. HTTP 200 or 206 scores 1, anything else 0. Multiplies the file score.",
           "A file nobody can download has no usable data, whatever else is true of it; the rule requires public access.",
           f"{RULE} (b)(3)"),
    "M2": (f"d = run date − `last_updated_on`. 1 if d ≤ {CYCLE_DAYS}, else e^(−(d−{CYCLE_DAYS})/{DECAY_DAYS}).",
           "The rule requires monthly updates. After the cycle the score falls fast (a week late 0.50, a month late 0.05) but never reaches 0, so very old files stay comparable.",
           f"{RULE} (b)(4)"),
    "M3": ("L = index `last_updated_on` − file `last_updated_on` (≤ 0 when the file is newer). Same curve as M2.",
           "An index much newer than its files describes plans whose rates may have changed (staleness).", "Quality"),
    "M4": ("Prices with a valid `expiration_date` on or after the run date ÷ prices.",
           "Expired rates are no longer payable (staleness). The share of `9999-12-31` (no end date) is reported.",
           f"{RULE} (b)(1)(i)(C)(2)"),
    "M5": ("Objects with every required field present, non-empty and correctly typed ÷ objects, averaged over object types (header, provider references, providers, TINs, items, rates, prices), per the schema version the file declares.",
           "Missing fields mean prices or providers cannot be read. Averaging over object types keeps a broken header from being drowned out by millions of prices.",
           "CMS schema (declared version)"),
    "M6": ("Values in the allowed set ÷ values checked (`negotiation_arrangement`, `billing_code_type`, `negotiated_type`, `billing_class`, `setting`, `tin.type`).",
           "Unknown values cannot be interpreted.", "CMS schema"),
    "M7": ("Objects meeting the condition ÷ objects it applies to: `service_code` on professional prices; `business_name` on EIN TINs; `plan_sponsor_name` on EIN plans.",
           "Without a place of service a professional rate cannot be matched to a setting.", "CMS schema"),
    "M8": ("Provider reference IDs used by prices that resolve to a provider group ÷ IDs used.",
           "A price that points to no provider cannot be used.", "CMS schema"),
    "M9": ("Items whose code fits its type (CPT 5 digits or 4 + F/T/U/M; HCPCS Level II on the CMS list, or CPT form; RC 4 digits; MS-DRG 3 digits; HIPPS 5; NDC 11) and have a description ÷ items.",
           "Malformed codes (revenue codes without leading zeros) break every join on the code: Medicare comparison, coverage and cross-payer benchmarks.",
           f"{RULE} (b)(1)(i)(B)"),
    "M10": ("NPIs that are 10 digits, start with 1 or 2 and pass the CMS Luhn check digit (constant 24) ÷ NPI entries; `0` (TIN-level contracts, allowed by the schema) is counted separately.",
            "An invalid NPI cannot be linked to a provider, address or specialty.", f"{RULE} (b)(1)(i)(C); CMS NPI spec"),
    "M11": ("TINs that are 9-digit EINs with an IRS-issued prefix (or a valid NPI when `tin.type = npi`) ÷ TIN entries.",
            "Invalid TINs break provider grouping; prefixes the IRS never issued, on individual clinicians, are consistent with SSNs, which the schema says to replace with the NPI.",
            f"{RULE} (b)(1)(i)(C); IRS EIN prefixes"),
    "M12": ("Place-of-service codes on the CMS list, or `CSTM-00`, ÷ codes.", "An unknown setting makes the rate ambiguous.",
            f"{RULE} (b)(1)(i)(C)"),
    "M13": ("Prices typed `negotiated`, `fee schedule`, `derived` or `per diem` ÷ prices (rate types).",
            "The rule requires dollar amounts unless none can be set in advance; percentages of charges cannot be compared.",
            f"{RULE} (b)(1)(i)(C)(1)"),
    "M14": ("1 − prices with `negotiated_rate` ≤ 0 ÷ prices.", "$0 or negative rates are placeholders, not prices; the schema requires a rate above 0.",
            "CMS schema; Quality"),
    "M15": (f"For professional CPT/HCPCS dollar prices with a 2026 Medicare fee (code, modifier, the file's state, facility or non-facility by place of service): x = ln(rate ÷ Medicare), pooled over the file. Score = 1 − share outside Q1 − {FENCE}·IQR … Q3 + {FENCE}·IQR; beyond {EXTREME}·IQR reported as extreme.",
            "Rates far from the file's usual ratio to Medicare are likely errors. Log scale because rates scale multiplicatively (skewness 1.13 raw, −0.16 log); IQR fences are data-driven rather than a fixed multiplier. Pooling via the ratio covers codes with few prices.",
            "Quality; CMS PFS RVU26D"),
    "M16": (f"Per price key with at least {MIN_PRICES} deduplicated prices: x = ln(rate). Score = 1 − share outside the same IQR fences.",
            "Catches data-entry and unit errors among prices for the same service.", "Quality"),
    "M17": ("Distinct NPIs found in NPPES and not deactivated (main file + monthly deactivation report) ÷ distinct NPIs.",
            "Rates for providers who can no longer bill are stale (provider validity).", "Quality"),
    "M18": ("1 − institutional prices on groups listing only individual (Type 1) NPIs ÷ institutional prices.",
            "Facility fees are paid to organizations; an individual-only list is likely wrong.", "Quality"),
    "M19": ("1 − (raw price rows − distinct group × service × rate rows) ÷ raw price rows.",
            "Duplicates inflate the file and every count taken from it (Health Net CA: 1.5 million rows, 39,960 distinct).", "Quality"),
    "M20": ("1 − billing codes appearing in two or more items ÷ distinct codes.",
            "Repeated items are the source of most duplicates and make it unclear which item is real.", "Quality"),
    "M21": ("Σ weight of the 583 HiLabs codes priced in the file (billing code or revenue-code sub-code) ÷ Σ weight of all 583; weight = sum of SPEC_PROC_WEIGHT over specialties.",
            "Code coverage weighted by what the business uses most: missing high-volume codes hurt negotiators most.", "Quality; HiLabs weights"),
    "M22": ("Active in-state NPIs with a HiLabs-specialty taxonomy in the file ÷ the same count in NPPES. **Reported, not scored.**",
            "Provider coverage measures how broad a network is, not whether the file is right: a narrow network published perfectly scores low (0.0003–0.37 here), so scoring it would penalize network design, not file quality.",
            "Quality"),
    "M23": ("Provider groups with a non-empty `network_name` ÷ groups.",
            "Without network names a file that mixes networks cannot be split by network (why C2b had to infer PPO rates). From March 2027 the rule requires one file per network.",
            "Quality; schema allows `[]`"),
    "M24": ("1 − provider (TIN + NPI) × service keys priced at different rates in groups whose `network_name` does not tell them apart ÷ keys priced in two or more groups.",
            "Different rates can be legitimate across networks, so only conflicts the file gives no way to resolve count. A negotiator cannot tell which rate applies.",
            "Quality"),
    "M25": ("File name matches the CMS pattern. **Not measurable here.**",
            "The CMS naming standard covers single-plan files only; files listed in an index have none, which every Centene rate file is.", "CMS guide"),
    "M26": ("`bundle` items with `bundled_codes` and `capitation` items with `covered_services` ÷ such items.",
            "A bundle or capitation price is meaningless without the services it covers.", "CMS guide (not enforced)"),
    "M27": ("Percentage prices given as whole percentages (1 < value ≤ 100) ÷ percentage prices.",
            "A value of 1.0 could mean 100% or 1%.", "CMS guide (not enforced)"),
    "I1": ("GET on the index URL; HTTP 200 or 206 scores 1, else 0. Multiplies the index score.",
           "An index nobody can download makes every file it lists unfindable.", f"{RULE} (b)(3)"),
    "I2": ("Same curve as M2 on the index's `last_updated_on`.", "Staleness of the plan list and links.", f"{RULE} (b)(4)"),
    "I3": ("Required fields present and valid, averaged over object types (root, reporting structures, plans, file locations).",
           "Missing fields leave plans or links unreadable.", "CMS table-of-contents schema"),
    "I4": ("`plan_id_type` ∈ {ein, hios} and `plan_market_type` ∈ {group, individual} ÷ values.", "Plans must be classifiable.", "CMS schema"),
    "I5": ("`plan_sponsor_name` on EIN plans; each reporting structure links an in-network or allowed-amount file.",
           "Sponsor and files must be identifiable.", "CMS schema"),
    "I6": ("Plan IDs matching their declared type (EIN: 9 digits with an IRS prefix; HIOS: 5 digits, or a 10/14-character ID with a valid state code) ÷ plans.",
           "Mislabelled IDs (HIOS IDs typed `ein`) break joins with CMS plan data and state detection.", f"CMS schema; {RULE} (b)(1)(i)(A)"),
    "I7": ("Plan IDs at a level the rule allows (14-character HIOS plan ID, 5-digit HIOS issuer ID, or EIN) ÷ plans; 10-character HIOS product IDs fail.",
           "A product ID covers many plans, so the index cannot say which plan a file serves; CMS lists 14-character plan IDs for 51 of Ambetter's 119 products.",
           f"{RULE} (b)(1)(i)(A)"),
    "I8": ("Reporting structures linking an allowed-amount (out-of-network) file ÷ structures.", "The rule requires the out-of-network file.",
           f"{RULE} (b)(1)(ii)"),
    "I9": ("Distinct linked in-network URLs answering HTTP 200/206 ÷ links.", "A dead link leaves its plans with no published rates.", f"{RULE} (b)(3)"),
    "I10": ("Same, for allowed-amount links.", "The required out-of-network file is effectively missing.", f"{RULE} (b)(3)"),
    "I11": ("Linked files whose top provider practice state equals the plans' state (HIOS state code, else C1 state) ÷ linked files. Provider states from the CMS Doctors and Clinicians file (Oct 2026), NPPES only where it has no record.",
            "Catches plans pointing to another state's network (Ambetter KS → a Kentucky network). NPPES addresses are often stale, so a fresher independent source comes first; a state over thousands of providers is robust.",
            "Quality"),
    "I12": ("Plan IDs that always carry the same type, market and issuer ÷ plan IDs.", "A plan must be identifiable.", "Quality"),
    "I13": ("1 − exact duplicate plan rows ÷ rows.", "Duplicates inflate plan counts.", "Quality"),
    "I14": ("Plans listed in only one reporting structure ÷ plans.", "Otherwise it is unclear which rate file applies.", "Quality"),
    "I15": ("HTTPS file locations ÷ locations.", "Plain HTTP downloads can be intercepted or altered.", "Quality"),
}
PS_TERMS = [("Code coverage", "M21"), ("Provider coverage", "M17 (scored), M22 (reported)"), ("$0 rates", "M14"),
            ("Outlier rates", "M15, M16"), ("Rate types", "M13, M27"), ("Duplicates", "M19, M20, I13"),
            ("Staleness", "M2, M3, M4, M17, I2")]


def read(name):
    with open(os.path.join(C3, name), encoding="utf-8") as f:
        return list(csv.DictReader(f))


def f3(x):
    return "–" if x in ("", None) else f"{float(x):.3f}"


def md_table(header, rows):
    out = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    out += ["| " + " | ".join(str(c).replace("|", "\\|") for c in r) + " |" for r in rows]
    return out


def main():
    argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter).parse_args()
    fm, fs, im, ix, ds, dg = (read(n) for n in ("c3_rate_file_metrics.csv", "c3_rate_file_scores.csv",
                                                "c3_index_metrics.csv", "c3_index_scores.csv",
                                                "c3_dataset_scores.csv", "c3_diagnosis.csv"))
    names = {r["metric_id"]: (r["metric"], r["group"]) for r in fm + im}
    order = [f"M{i}" for i in range(1, 28)] + [f"I{i}" for i in range(1, 16)]
    defs = [(m, names[m][0], GROUP_NAME[names[m][1]], *DEF[m]) for m in order]
    ds_sorted = sorted(ds, key=lambda r: -float(r["dataset_score"]))
    fs_sorted = sorted(fs, key=lambda r: (r["dataset"], -float(r["file_score"])))
    hn = {r["metric_id"]: r for r in fm if r["rate_file"] == "healthnet-ca"}

    md = ["# File quality metrics (deliverable 3)", "",
          "The metrics used to grade the price transparency (PTD) files, how each is calculated, and why, with the "
          "score of every file we use. Generated by `pipeline/c3_report.py` from the C3 outputs; the full reasoning "
          "and decision history are in `LOGIC.md` section 3C.", "",
          "## 1. What is graded", "",
          "- **Every Centene file:** the 4 index files (tables of contents) and the 33 rate files they link: Ambetter 28, "
          "Health Net 2, Fidelis 2, QualChoice 1. Our in-scope file (Health Net CA, commercial group PPO) is one of them; "
          "the rest show the metrics work on any payer's file.",
          "- **Against what:** 45 CFR §147.212 as in force on 1 Oct 2025 (the October 2026 rule keeps it until 6 Mar 2027), "
          "the CMS JSON schema at the version each file declares (v2.0.0; Ambetter OH v1.0.0), and data-quality checks "
          "against CMS reference data (2026 Medicare fee schedule, HCPCS, place-of-service codes, NPPES, Doctors and "
          "Clinicians) and the HiLabs code weights.", "",
          "## 2. How the scores are built", ""]
    md += md_table(["Step", "Calculation", "Why"], [
        ["Metric score", "Each metric gives numerator, denominator, value and a score from 0 to 1 (higher is better); "
         "\"share of bad\" metrics score 1 − share.", "One scale for every metric; raw counts stay visible."],
        ["Gate", "Reachable (M1, I1): 1 or 0, multiplies the file or index score.", "A file nobody can download has no usable data."],
        ["Group score", "Plain average of the group's measured metrics.",
         "Multiplying would let one weak metric wipe out a group; averaging gives partial credit and is easy to explain."],
        ["Rate-file score", "Gate × weighted average of groups: " + ", ".join(
            f"{GROUP_NAME[g]} {w:.0%}" for g, w in FILE_GROUPS.items()) + ".",
         "Compliance (what the rule and schema require) weighs most; then plausibility, because wrong values mislead negotiators."],
        ["Index score", "Gate × weighted average of groups: " + ", ".join(
            f"{GROUP_NAME[g]} {w:.0%}" for g, w in INDEX_GROUPS.items()) + ".",
         "An index holds no prices, so plausibility is replaced by consistency with the files it links."],
        ["Unmeasurable metrics", "Left out, weights re-normalized over what was measured; listed with the reason.",
         "Scoring them 0 would punish the file for our missing reference data; 1 would reward it."],
        ["Requirements vs guidance", "Compliance holds only what the rule or schema enforces; CMS README guidance sits in usability, marked (guide).",
         "A file is not called non-compliant for breaking an unenforced convention."],
        ["Dataset score", f"index score^{INDEX_SHARE} × (mean rate-file score)^{1 - INDEX_SHARE:.1f}; an unreachable rate file counts 0.",
         "Index and rate files are links in one chain, so their qualities multiply and a zero in either gives 0; "
         "the exponents give the rate files (where the prices are) more influence and keep the 0–1 scale."],
    ])
    md += ["", "**Problem-statement terms and the metrics that cover them:**", ""]
    md += md_table(["Term", "Metrics"], PS_TERMS)
    for title, prefix in (("3. Rate-file metrics", "M"), ("4. Index-file metrics", "I")):
        md += ["", f"## {title}", ""]
        md += md_table(["#", "Metric", "Group", "How it is calculated", "Why", "Basis"],
                       [d for d in defs if d[0].startswith(prefix)])

    md += ["", "## 5. Scores", "", "### Datasets", ""]
    md += md_table(["Dataset", "Index score", "Rate files (reachable)", "Mean rate-file score", "Lowest file", "**Dataset score**"],
                   [[DISPLAY[r["dataset"]], f3(r["index_score"]), f"{r['rate_files']} ({r['rate_files_reachable']})",
                     f3(r["mean_rate_file_score"]), f3(r["min_rate_file_score"]), f"**{f3(r['dataset_score'])}**"]
                    for r in ds_sorted])
    md += ["", "### Index files", ""]
    md += md_table(["Index", "Gate", *[GROUP_NAME[g] for g in INDEX_GROUPS], "Index score"],
                   [[r["index"], f3(r["gate"]), *[f3(r[g]) for g in INDEX_GROUPS], f"**{f3(r['index_score'])}**"]
                    for r in sorted(ix, key=lambda r: -float(r["index_score"]))])
    md += ["", "### Rate files (all 33)", ""]
    md += md_table(["Rate file", "Dataset", "State", "Gate", *[GROUP_NAME[g] for g in FILE_GROUPS], "File score"],
                   [[r["rate_file"], DISPLAY[r["dataset"]], r["state"], f3(r["gate"]),
                     *[f3(r[g]) for g in FILE_GROUPS], f"**{f3(r['file_score'])}**"] for r in fs_sorted])
    md += ["", "### In-scope file in detail: Health Net CA", ""]
    md += md_table(["#", "Metric", "Numerator", "Denominator", "Score", "Note"],
                   [[m, names[m][0], hn[m]["numerator"], hn[m]["denominator"], f3(hn[m]["score"]), hn[m]["note"]]
                    for m in order if m.startswith("M")])

    md += ["", "## 6. What drove the scores", "",
           "Problems ranked by the dataset points (0–100) they cost; severity High ≥ 2, Medium ≥ 0.5, Low below, with "
           "an unreachable file Critical and a wrong-state link at least High. Full list with evidence and fixes: "
           "`output/c3/c3_diagnosis.md` (`pipeline/c3_diagnose.py`).", ""]
    rows = [r for r in dg if r["severity"] in ("Critical", "High", "Medium")]
    md += md_table(["Dataset", "Severity", "Metric", "Points", "Problem", "Fix (publisher)"],
                   [[DISPLAY[r["dataset"]], r["severity"], f"{r['metric_id']} {r['metric']}", r["points_lost"],
                     r["problem"], r["fix_publisher"]] for r in rows])
    md += ["", "## 7. Limits", "",
           "- Thirteen checks are at or above 0.99 for every reachable Centene file, nine at exactly 1.0: the files meet "
           "the enforced schema, so compliance separates them little. The checks still matter for other payers.",
           "- M15 fences are relative to each file's own spread of rate ÷ Medicare: it measures consistency, not level; "
           "Medicare amounts use state-average GPCIs, not localities.",
           "- I11 uses Medicare-enrolled clinicians first (76% of NPIs in the KS file); organization NPIs rely on NPPES.",
           "- I7 accepts 5-digit issuer IDs because we cannot tell whether a 14-character plan ID was available.",
           "- Weights and the dataset exponents are judgements, stated and fixed in `LOGIC.md` (R3.5, R3.11, R3.14).", ""]
    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, "03_file_quality_metrics.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(md).replace("$", "\\$"))  # GitHub renders text between two $ as math

    # ---- workbook
    wb = Workbook()
    sheets = [
        ("Metric definitions", ["Metric", "Name", "Group", "How it is calculated", "Why", "Basis"],
         [[c.replace("**", "") for c in d] for d in defs]),
        ("Dataset scores", ["Dataset", "Index score", "Rate files", "Reachable", "Mean rate-file score",
                            "Lowest rate-file score", "Dataset score"],
         [[DISPLAY[r["dataset"]], float(r["index_score"]), int(r["rate_files"]), int(r["rate_files_reachable"]),
           float(r["mean_rate_file_score"]), float(r["min_rate_file_score"] or 0), float(r["dataset_score"])] for r in ds_sorted]),
        ("Index scores", ["Index", "Dataset", "Gate", *[GROUP_NAME[g] for g in INDEX_GROUPS], "Index score"],
         [[r["index"], DISPLAY[r["dataset"]], float(r["gate"]),
           *[float(r[g]) if r[g] else None for g in INDEX_GROUPS], float(r["index_score"])] for r in ix]),
        ("Rate file scores", ["Rate file", "Dataset", "State", "Gate", *[GROUP_NAME[g] for g in FILE_GROUPS],
                              "File score", "Metrics measured"],
         [[r["rate_file"], DISPLAY[r["dataset"]], r["state"], float(r["gate"]),
           *[float(r[g]) if r[g] else None for g in FILE_GROUPS], float(r["file_score"]), int(r["metrics_measured"])]
          for r in fs_sorted]),
        ("Rate file metrics", ["Rate file", "Dataset", "Metric", "Name", "Group", "Numerator", "Denominator", "Value",
                               "Score", "Measured", "Note"],
         [[r["rate_file"], DISPLAY[r["dataset"]], r["metric_id"], r["metric"], GROUP_NAME[r["group"]],
           r["numerator"], r["denominator"], r["value"], float(r["score"]) if r["score"] else None,
           r["measured"], r["note"]] for r in fm]),
        ("Index metrics", ["Index", "Dataset", "Metric", "Name", "Group", "Numerator", "Denominator", "Value", "Score",
                           "Measured", "Note"],
         [[r["index"], DISPLAY[r["dataset"]], r["metric_id"], r["metric"], GROUP_NAME[r["group"]], r["numerator"],
           r["denominator"], r["value"], float(r["score"]) if r["score"] else None, r["measured"], r["note"]] for r in im]),
        ("Problems", ["Dataset", "Rank", "Severity", "Metric", "Name", "Points lost", "Where", "Numbers", "Evidence",
                      "Problem", "Why", "Fix (publisher)", "Fix (our pipeline)"],
         [[DISPLAY[r["dataset"]], r["rank"], r["severity"], r["metric_id"], r["metric"], r["points_lost"], r["where"],
           r["numbers"], r["evidence"], r["problem"], r["why"], r["fix_publisher"], r["fix_pipeline"]] for r in dg]),
    ]
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    for i, (title, header, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), head_fill
        ws.freeze_panes = "A2"
        for j, h in enumerate(header, 1):
            longest = max([len(str(h))] + [len(str(r[j - 1] or "")) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, longest + 2), 70)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "0.000"
                elif isinstance(c.value, str) and len(c.value) > 70:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(os.path.join(OUT, "03_file_quality_scores.xlsx"))
    print(f"wrote {os.path.join(OUT, '03_file_quality_metrics.md')} and 03_file_quality_scores.xlsx")


if __name__ == "__main__":
    main()
