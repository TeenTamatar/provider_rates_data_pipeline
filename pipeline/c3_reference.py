"""C3 reference tables: the outside lists the rate-file metrics are checked against.

Builds small Parquet tables in output/c3/ref/ so the metric step never rescans the big sources:

  nppes.parquet           npi, entity_type, practice_state, deactivated (NPPES main file + deactivation report),
                          taxonomies (list of up to 15 codes)                    -> M17, M18, M22
  hilabs_codes.parquet    code, weight (sum of SPEC_PROC_WEIGHT over specialties) -> M21
  hilabs_taxonomies.parquet taxonomy (Taxonomy-Spec Xwalk)                        -> M22
  hcpcs.parquet           code (CMS Alpha-Numeric HCPCS file, Level II)           -> M9
  medicare_pfs.parquet    code, modifier, state, nonfacility, facility (2026 PFS, state-average GPCI) -> M15
  provider_states.parquet npi, state, src: every state where a provider practises, by source        -> I11
                          dac (CMS Doctors and Clinicians, all locations), nppes_primary, nppes_secondary

Sources (see LOGIC.md 3C): NPPES October 2026, NPPES deactivation report 2026-10-12, HiLabs workbook,
CMS HCPCS October 2026, CMS PFS RVU26D (October 2026 release), CMS Doctors and Clinicians (2026-10-02).
The two CMS zip files are unzipped into Data/cms_ref/rvu26d/ and Data/cms_ref/hcpcs_2026_10/:
  https://www.cms.gov/files/zip/rvu26d-updated-08-26-2026.zip
  https://www.cms.gov/files/zip/october-2026-alpha-numeric-hcpcs-file.zip
The Doctors and Clinicians file is expected at Data/c5_truth/DAC_NationalDownloadableFile.csv.

Order: c3_reference.py -> rate_scan.py -> c3_metrics.py -> c3_index_metrics.py
(rate_scan.py and c3_metrics.py also need output/c1/c1_rate_files.csv from C1, and every rate file in
Data/rate_files/: pipeline/download_rate_files.py).

Usage:
    python pipeline/c3_reference.py
"""
import argparse
import glob
import json
import os
import time

import duckdb
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "Data")
DEFAULT_OUT = os.path.join(ROOT, "output", "c3", "ref")


def build_nppes(con, out):
    src = glob.glob(os.path.join(DATA, "NPPES_Data_Dissemination_*", "npidata_pfile_*[0-9].csv"))[0]
    tax = ", ".join(f'"Healthcare Provider Taxonomy Code_{i}"' for i in range(1, 16))
    con.execute(f"""
        create table nppes as
        select cast("NPI" as bigint) npi,
               "Entity Type Code" entity_type,
               "Provider Business Practice Location Address State Name" practice_state,
               ("NPI Deactivation Date" is not null and "NPI Deactivation Date" <> ''
                and ("NPI Reactivation Date" is null or "NPI Reactivation Date" = '')) deactivated,
               list_filter([{tax}], x -> x is not null and x <> '') taxonomies
        from read_csv('{src}', header=true, all_varchar=true)""")
    # The monthly deactivation report is newer than the main file's flags.
    wb = openpyxl.load_workbook(glob.glob(os.path.join(DATA, "NPPES Deactivated NPI Report*.xlsx"))[0], read_only=True)
    deact = [int(r[0]) for r in wb.active.iter_rows(min_row=3, values_only=True) if r[0] and str(r[0]).isdigit()]
    con.execute("create table deact(npi bigint)")
    con.executemany("insert into deact values (?)", [(n,) for n in deact])
    con.execute("""create table nppes2 as
                   select n.npi, n.entity_type, n.practice_state,
                          n.deactivated or d.npi is not null deactivated, n.taxonomies
                   from nppes n left join (select distinct npi from deact) d using (npi)""")
    con.execute(f"copy nppes2 to '{os.path.join(out, 'nppes.parquet')}' (format parquet)")
    return {"nppes_rows": con.execute("select count(*) from nppes2").fetchone()[0],
            "deactivation_report_npis": len(deact),
            "deactivated_total": con.execute("select count(*) from nppes2 where deactivated").fetchone()[0]}


def build_hilabs(con, out):
    wb = openpyxl.load_workbook(os.path.join(DATA, "Rate Benchmarking Individual Weights_Hackathon.xlsx"), read_only=True)
    rows = list(wb["Spec-Proc Weights"].iter_rows(min_row=2, values_only=True))
    weights = {}
    for _, _, code, _, w in rows:
        if code:
            weights[str(code).strip()] = weights.get(str(code).strip(), 0) + float(w or 0)
    con.execute("create table hc(code varchar, weight double)")
    con.executemany("insert into hc values (?, ?)", list(weights.items()))
    con.execute(f"copy hc to '{os.path.join(out, 'hilabs_codes.parquet')}' (format parquet)")
    taxa = {str(r[3]).strip() for r in wb["Taxonomy-Spec Xwalk"].iter_rows(min_row=2, values_only=True) if r[3]}
    con.execute("create table ht(taxonomy varchar)")
    con.executemany("insert into ht values (?)", [(t,) for t in sorted(taxa)])
    con.execute(f"copy ht to '{os.path.join(out, 'hilabs_taxonomies.parquet')}' (format parquet)")
    return {"hilabs_codes": len(weights), "hilabs_taxonomies": len(taxa)}


def build_hcpcs(con, out):
    # Fixed-width file: the code is the first 5 characters (record layout HCPC2026_recordlayout.txt).
    src = glob.glob(os.path.join(DATA, "cms_ref", "hcpcs_*", "HCPC*_ANWEB_*.txt"))[0]
    codes = set()
    with open(src, encoding="latin-1") as f:
        for line in f:
            c = line[:5].strip()
            if len(c) == 5:
                codes.add(c)
    con.execute("create table hx(code varchar)")
    con.executemany("insert into hx values (?)", [(c,) for c in sorted(codes)])
    con.execute(f"copy hx to '{os.path.join(out, 'hcpcs.parquet')}' (format parquet)")
    return {"hcpcs_level2_codes": len(codes), "hcpcs_source": os.path.basename(src)}


def build_pfs(con, out):
    d = glob.glob(os.path.join(DATA, "cms_ref", "rvu26*"))
    d = [p for p in d if os.path.isdir(p)][0]
    rvu = glob.glob(os.path.join(d, "PPRRVU*_nonQPP.csv"))[0]
    gpci = glob.glob(os.path.join(d, "GPCI*.csv"))[0]
    # PPRRVU: 9 header lines before the column row "HCPCS,MOD,DESCRIPTION,...".
    con.execute(f"""create table rvu as select * from read_csv('{rvu}', skip=9, header=true, all_varchar=true,
                    null_padding=true, ignore_errors=true)""")
    cols = [r[0] for r in con.execute("describe rvu").fetchall()]
    # Columns by position (names repeat in the multi-line header): 0 HCPCS, 1 MOD, 3 STATUS, 5 WORK RVU,
    # 6 NON-FAC PE RVU, 8 FACILITY PE RVU, 10 MP RVU, 25 CONV FACTOR.
    c = lambda i: f'"{cols[i]}"'
    con.execute(f"""create table g as select "State" state, try_cast("2026 PW GPCI (with 1.0 Floor)***" as double) pw,
                    try_cast("2026 PE GPCI" as double) pe, try_cast("2026 MP GPCI" as double) mp
                    from read_csv('{gpci}', skip=2, header=true, all_varchar=true) where "State" is not null""")
    con.execute(f"""create table pfs as
        with r as (select {c(0)} code, coalesce({c(1)}, '') modifier, {c(3)} status,
                          try_cast({c(5)} as double) rvu_work, try_cast({c(6)} as double) rvu_pe_nf,
                          try_cast({c(8)} as double) rvu_pe_f, try_cast({c(10)} as double) rvu_mp,
                          try_cast({c(25)} as double) cf
                   from rvu where {c(3)} in ('A', 'R', 'T')),
             s as (select state, avg(pw) pw, avg(pe) pe, avg(mp) mp from g group by state)
        select r.code, r.modifier, s.state,
               round((r.rvu_work * s.pw + r.rvu_pe_nf * s.pe + r.rvu_mp * s.mp) * r.cf, 2) nonfacility,
               round((r.rvu_work * s.pw + r.rvu_pe_f * s.pe + r.rvu_mp * s.mp) * r.cf, 2) facility
        from r cross join s
        where (r.rvu_work * s.pw + r.rvu_pe_nf * s.pe + r.rvu_mp * s.mp) * r.cf > 0""")
    con.execute(f"copy pfs to '{os.path.join(out, 'medicare_pfs.parquet')}' (format parquet)")
    return {"pfs_rows": con.execute("select count(*) from pfs").fetchone()[0],
            "pfs_codes": con.execute("select count(distinct code) from pfs").fetchone()[0],
            "pfs_states": con.execute("select count(distinct state) from pfs").fetchone()[0],
            "conversion_factor": con.execute(f"select distinct try_cast({c(25)} as double) from rvu where {c(3)} = 'A'").fetchall()}


def build_provider_states(con, out):
    """All practice states per NPI from three sources. NPPES addresses are often stale, so I11 prefers the
    Medicare file (independent, refreshed monthly) and uses NPPES only for providers it does not cover."""
    dac = glob.glob(os.path.join(DATA, "c5_truth", "DAC_NationalDownloadableFile.csv"))[0]
    pl = glob.glob(os.path.join(DATA, "NPPES_Data_Dissemination_*", "pl_pfile_*[0-9].csv"))[0]
    con.execute(f"""create table ps as
        select distinct try_cast("NPI" as bigint) npi, upper(trim("State")) state, 'dac' src
        from read_csv('{dac}', header=true, all_varchar=true) where "State" is not null
        union all
        select npi, upper(trim(practice_state)), 'nppes_primary' from '{out}/nppes.parquet'
        where not deactivated and practice_state is not null
        union all
        select distinct try_cast("NPI" as bigint), upper(trim("Provider Secondary Practice Location Address - State Name")),
               'nppes_secondary'
        from read_csv('{pl}', header=true, all_varchar=true)
        where "Provider Secondary Practice Location Address - State Name" is not null""")
    con.execute(f"copy (select * from ps where npi is not null) to '{os.path.join(out, 'provider_states.parquet')}' (format parquet)")
    return {"provider_states_by_source": dict(con.execute("select src, count(distinct npi) from ps group by 1").fetchall())}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=DEFAULT_OUT)
    ap.add_argument("--skip-nppes", action="store_true", help="keep an existing nppes.parquet")
    ap.add_argument("--only-states", action="store_true", help="rebuild provider_states.parquet only")
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    con = duckdb.connect()
    t0 = time.time()
    summary = {}
    if args.only_states:
        print(json.dumps(build_provider_states(con, args.out), indent=2))
        return
    summary.update(build_hilabs(con, args.out))
    summary.update(build_hcpcs(con, args.out))
    summary.update(build_pfs(con, args.out))
    if not args.skip_nppes:
        summary.update(build_nppes(con, args.out))
    summary.update(build_provider_states(con, args.out))
    summary["seconds"] = round(time.time() - t0, 1)
    with open(os.path.join(args.out, "c3_reference_summary.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2, default=str)
    print(json.dumps(summary, indent=2, default=str))


if __name__ == "__main__":
    main()
